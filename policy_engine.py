import re
import pandas as pd

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

from database import initialize_database, save_run


# ============================================================
# PATHS
# ============================================================

BASE = Path(__file__).resolve().parent

INPUT = (
    BASE
    / "generated_data"
    / "synthetic_data.xlsx"
)

OUTPUT = BASE / "generated_data" / "policy_results.xlsx"


# ============================================================
# RULE DESCRIPTIONS AND REMEDIATION
# ============================================================

RULES = {

    "PII-01": (
        "Full name detected",
        "Redact or tokenise customer_name"
    ),

    "PII-02": (
        "Personal email detected",
        "Mask email or replace with token"
    ),

    "PII-03": (
        "Phone number detected",
        "Mask phone number"
    ),

    "PII-04": (
        "Postal address detected",
        "Remove or generalise address"
    ),

    "PII-05": (
        "National Insurance number detected",
        "Remove ni_number"
    ),

    "PII-06": (
        "Passport number detected",
        "Remove passport_number"
    ),

    "PII-08": (
        "Credit card number detected",
        "Remove or mask card number"
    ),

    "PII-09": (
        "IP address detected",
        "Anonymise IP address"
    ),

    "SPII-01": (
        "Health or medical information detected",
        "Remove sensitive field or exclude record"
    ),

    "SPII-02": (
        "Ethnicity detected",
        "Remove sensitive field or exclude record"
    ),

    "SPII-03": (
        "Religion detected",
        "Remove sensitive field or exclude record"
    ),

    "SPII-04": (
        "Political opinion detected",
        "Remove sensitive field or exclude record"
    ),

    "CPII-01": (
        "Name plus date of birth combination detected",
        "Remove name or generalise DOB"
    ),

    "CPII-02": (
        "Name plus address combination detected",
        "Remove name or generalise address"
    ),

    "CPII-03": (
        "Name plus phone combination detected",
        "Redact name and mask phone"
    ),

    "CPII-04": (
        "Name plus email combination detected",
        "Redact name and mask email"
    ),

    "CPII-05": (
        "Date of birth plus gender combination detected",
        "Generalise DOB or remove gender"
    ),

    "CPII-06": (
        "Employee identifier plus department plus role combination detected",
        "Tokenise employee_id and generalise role/department"
    ),

    "CPII-08": (
        "Free-text comments contain personal identifier",
        "Redact PII from feedback"
    )
}


# ============================================================
# RULES THAT CAUSE BLOCK
# ============================================================

BLOCK_RULES = {

    "PII-05",
    "PII-06",

    "SPII-01",
    "SPII-02",
    "SPII-03",
    "SPII-04"
}


# ============================================================
# CHECK ONE RECORD
# ============================================================

def check_record(row):

    triggered = []

    def add(rule):

        if rule not in triggered:
            triggered.append(rule)

    # --------------------------------------------------------
    # Direct PII
    # --------------------------------------------------------

    if (
        pd.notna(row["customer_name"])
        and str(row["customer_name"]).strip()
    ):
        add("PII-01")

    if (
        pd.notna(row["email"])
        and str(row["email"]).strip()
    ):
        add("PII-02")

    if (
        pd.notna(row["phone"])
        and str(row["phone"]).strip()
    ):
        add("PII-03")

    if (
        pd.notna(row["address"])
        and str(row["address"]).strip()
    ):
        add("PII-04")

    if (
        pd.notna(row["ni_number"])
        and str(row["ni_number"]).strip()
    ):
        add("PII-05")

    if (
        pd.notna(row["passport_number"])
        and str(row["passport_number"]).strip()
    ):
        add("PII-06")

    if (
        pd.notna(row["credit_card_number"])
        and str(row["credit_card_number"]).strip()
    ):
        add("PII-08")

    if (
        pd.notna(row["ip_address"])
        and str(row["ip_address"]).strip()
    ):
        add("PII-09")


    # --------------------------------------------------------
    # Sensitive PII
    # --------------------------------------------------------

    if (
        pd.notna(row["medical_condition"])
        and str(row["medical_condition"]).strip()
    ):
        add("SPII-01")

    if (
        pd.notna(row["ethnicity"])
        and str(row["ethnicity"]).strip()
    ):
        add("SPII-02")

    if (
        pd.notna(row["religion"])
        and str(row["religion"]).strip()
    ):
        add("SPII-03")

    if (
        pd.notna(row["political_view"])
        and str(row["political_view"]).strip()
    ):
        add("SPII-04")


    # --------------------------------------------------------
    # Combination rules
    # --------------------------------------------------------

    if (
        pd.notna(row["customer_name"])
        and pd.notna(row["dob"])
    ):
        add("CPII-01")

    if (
        pd.notna(row["customer_name"])
        and pd.notna(row["address"])
    ):
        add("CPII-02")

    if (
        pd.notna(row["customer_name"])
        and pd.notna(row["phone"])
    ):
        add("CPII-03")

    if (
        pd.notna(row["customer_name"])
        and pd.notna(row["email"])
    ):
        add("CPII-04")

    if (
        pd.notna(row["dob"])
        and pd.notna(row["gender"])
    ):
        add("CPII-05")

    if (
        pd.notna(row["employee_id"])
        and pd.notna(row["department"])
        and pd.notna(row["job_role"])
    ):
        add("CPII-06")


    # --------------------------------------------------------
    # Free-text PII
    # --------------------------------------------------------

    feedback = (
        str(row["feedback"])
        if pd.notna(row["feedback"])
        else ""
    )

    if re.search(
        r"[\w\.-]+@[\w\.-]+\.\w+|07\d{9}|\b\d{10}\b",
        feedback
    ):
        add("CPII-08")


    # --------------------------------------------------------
    # Determine outcome
    # --------------------------------------------------------

    if not triggered:

        outcome = "PASS"

        reason = "No PII detected"

        remediation = "No action required"

    elif any(
        rule in BLOCK_RULES
        for rule in triggered
    ):

        outcome = "BLOCK"

        reason = "; ".join(
            f"{rule}: {RULES[rule][0]}"
            for rule in triggered
        )

        remediation = "; ".join(
            dict.fromkeys(
                RULES[rule][1]
                for rule in triggered
            )
        )

    else:

        outcome = "FLAG"

        reason = "; ".join(
            f"{rule}: {RULES[rule][0]}"
            for rule in triggered
        )

        remediation = "; ".join(
            dict.fromkeys(
                RULES[rule][1]
                for rule in triggered
            )
        )

    return [
        outcome,
        "; ".join(triggered),
        reason,
        remediation
    ]


# ============================================================
# INITIALIZE DATABASE
# ============================================================

initialize_database()


# ============================================================
# READ SYNTHETIC DATA
# ============================================================

df = pd.read_excel(INPUT)


# ============================================================
# RUN POLICY ENGINE
# ============================================================

results = df.apply(
    check_record,
    axis=1,
    result_type="expand"
)

results.columns = [
    "expected_outcome",
    "expected_rule_triggers",
    "expected_reason",
    "suggested_remediation"
]


# ============================================================
# ADD POLICY RESULTS
# ============================================================

df = pd.concat(
    [df, results],
    axis=1
)


# ============================================================
# SAVE COMPLETE RUN TO DATABASE
# ============================================================

records = df.to_dict(
    orient="records"
)

run_id = save_run(
    records=records,
    dataset_name=INPUT.name,
    policy_name="AI Training Data PII Policy"
)

print(
    f"\nSaved evaluation to database as Run #{run_id}"
)


# ============================================================
# SAVE RESULTS TO EXCEL
# ============================================================

OUTPUT.parent.mkdir(
    parents=True,
    exist_ok=True
)

df.to_excel(
    OUTPUT,
    index=False,
    sheet_name="Policy Results"
)


# ============================================================
# FORMAT EXCEL
# ============================================================

wb = load_workbook(OUTPUT)

ws = wb["Policy Results"]

ws.freeze_panes = "A2"

ws.auto_filter.ref = ws.dimensions


# ------------------------------------------------------------
# Header formatting
# ------------------------------------------------------------

for cell in ws[1]:

    cell.font = Font(
        bold=True
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# ------------------------------------------------------------
# Column widths
# ------------------------------------------------------------

widths = {

    "A": 12,
    "B": 22,
    "C": 30,
    "D": 18,
    "E": 35,
    "F": 15,
    "G": 12,
    "H": 20,
    "I": 18,
    "J": 22,
    "K": 18,
    "L": 25,
    "M": 16,
    "N": 16,
    "O": 20,
    "P": 16,
    "Q": 16,
    "R": 18,
    "S": 22,
    "T": 18,
    "U": 45,
    "V": 18,
    "W": 35,
    "X": 70,
    "Y": 70
}

for column, width in widths.items():

    ws.column_dimensions[column].width = width


# ------------------------------------------------------------
# Wrap all cells
# ------------------------------------------------------------

for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


ws.row_dimensions[1].height = 35

wb.save(OUTPUT)


# ============================================================
# PRINT SUMMARY
# ============================================================

print("\nPolicy evaluation complete!")

print(f"Output: {OUTPUT}")

print(f"Run ID: {run_id}")

print("\nOutcome summary:")

print(
    df["expected_outcome"]
    .value_counts()
)

print("\nRule summary:")

print(
    df["expected_rule_triggers"]
    .fillna("")
    .str.split("; ")
    .explode()
    .replace("", pd.NA)
    .dropna()
    .value_counts()
)