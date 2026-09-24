import random
import pandas as pd
from faker import Faker
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

fake = Faker()

N = 100

CUSTOMERS = [f"CUS{i:03d}" for i in range(1, 6)]

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "generated_data"
OUTPUT_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = OUTPUT_DIR / "synthetic_data.xlsx"


# ============================================================
# COLUMNS
# ============================================================

COLUMNS = [
    "record_id",
    "customer_name",
    "email",
    "phone",
    "address",
    "dob",
    "gender",
    "passport_number",
    "ni_number",
    "credit_card_number",
    "bank_account",
    "medical_condition",
    "ethnicity",
    "religion",
    "political_view",
    "employee_id",
    "department",
    "job_role",
    "customer_id",
    "ip_address",
    "feedback"
]


# ============================================================
# POLICY SCENARIOS
# ============================================================

SCENARIOS = [
    "PASS",

    # Direct PII
    "PII-01",
    "PII-02",
    "PII-03",
    "PII-04",
    "PII-05",
    "PII-06",
    "PII-08",
    "PII-09",

    # Sensitive PII
    "SPII-01",
    "SPII-02",
    "SPII-03",
    "SPII-04",

    # Combined PII
    "CPII-01",
    "CPII-02",
    "CPII-03",
    "CPII-04",
    "CPII-05",
    "CPII-06",
    "CPII-08",

    # Other scenarios
    "FREE_TEXT",
    "MULTIPLE_DIRECT",
    "MULTIPLE_SENSITIVE",
    "MIXED"
]


# ============================================================
# SCENARIO GROUPS
# ============================================================

# These scenarios should result in PASS.
PASS_SCENARIOS = [
    "PASS"
]


# These scenarios contain sensitive information that
# should normally result in BLOCK according to the policy.
BLOCK_SCENARIOS = [
    "PII-05",
    "PII-06",
    "SPII-01",
    "SPII-02",
    "SPII-03",
    "SPII-04",
    "MULTIPLE_SENSITIVE",
    "MIXED"
]


# These scenarios normally result in FLAG.
FLAG_SCENARIOS = [
    "PII-01",
    "PII-02",
    "PII-03",
    "PII-04",
    "PII-08",
    "PII-09",
    "CPII-01",
    "CPII-02",
    "CPII-03",
    "CPII-04",
    "CPII-05",
    "CPII-06",
    "CPII-08",
    "FREE_TEXT",
    "MULTIPLE_DIRECT"
]


# ============================================================
# CREATE EMPTY RECORD
# ============================================================

def empty_record(i):

    return {
        "record_id": i,

        "customer_name": None,
        "email": None,
        "phone": None,
        "address": None,
        "dob": None,
        "gender": None,

        "passport_number": None,
        "ni_number": None,
        "credit_card_number": None,
        "bank_account": None,

        "medical_condition": None,
        "ethnicity": None,
        "religion": None,
        "political_view": None,

        "employee_id": None,
        "department": None,
        "job_role": None,

        "customer_id": (
            random.choice(CUSTOMERS)
            if random.random() < 0.2
            else None
        ),

        "ip_address": None,

        "feedback": random.choice([
            "Good service",
            "Customer satisfied",
            "Transaction completed successfully",
            "Mobile application works well"
        ])
    }


# ============================================================
# GENERATE ONE SCENARIO
# ============================================================

def generate_record(i, scenario):

    r = empty_record(i)

    # --------------------------------------------------------
    # PASS
    # --------------------------------------------------------

    if scenario == "PASS":
        pass


    # --------------------------------------------------------
    # DIRECT PII
    # --------------------------------------------------------

    elif scenario == "PII-01":

        r["customer_name"] = fake.name()


    elif scenario == "PII-02":

        r["email"] = fake.email()


    elif scenario == "PII-03":

        r["phone"] = fake.numerify("07#########")


    elif scenario == "PII-04":

        r["address"] = fake.address().replace("\n", ", ")


    elif scenario == "PII-05":

        r["ni_number"] = fake.bothify(
            "??######?"
        ).upper()


    elif scenario == "PII-06":

        r["passport_number"] = fake.numerify(
            "#########"
        )


    elif scenario == "PII-08":

        r["credit_card_number"] = fake.credit_card_number()


    elif scenario == "PII-09":

        r["ip_address"] = fake.ipv4()


    # --------------------------------------------------------
    # SENSITIVE PII
    # --------------------------------------------------------

    elif scenario == "SPII-01":

        r["medical_condition"] = random.choice([
            "Diabetes",
            "Asthma",
            "Hypertension"
        ])


    elif scenario == "SPII-02":

        r["ethnicity"] = random.choice([
            "Asian",
            "Black",
            "White",
            "Mixed"
        ])


    elif scenario == "SPII-03":

        r["religion"] = random.choice([
            "Christian",
            "Hindu",
            "Muslim",
            "Sikh"
        ])


    elif scenario == "SPII-04":

        r["political_view"] = random.choice([
            "Conservative",
            "Labour",
            "Liberal"
        ])


    # --------------------------------------------------------
    # COMBINED PII
    # --------------------------------------------------------

    elif scenario == "CPII-01":

        r["customer_name"] = fake.name()

        r["dob"] = fake.date_of_birth(
            minimum_age=18,
            maximum_age=80
        ).isoformat()


    elif scenario == "CPII-02":

        r["customer_name"] = fake.name()

        r["address"] = fake.address().replace(
            "\n",
            ", "
        )


    elif scenario == "CPII-03":

        r["customer_name"] = fake.name()

        r["phone"] = fake.numerify(
            "07#########"
        )


    elif scenario == "CPII-04":

        r["customer_name"] = fake.name()

        r["email"] = fake.email()


    elif scenario == "CPII-05":

        r["dob"] = fake.date_of_birth(
            minimum_age=18,
            maximum_age=80
        ).isoformat()

        r["gender"] = random.choice([
            "Male",
            "Female"
        ])


    elif scenario == "CPII-06":

        r["employee_id"] = (
            f"EMP{random.randint(1000, 9999)}"
        )

        r["department"] = random.choice([
            "Finance",
            "IT",
            "HR",
            "Operations"
        ])

        r["job_role"] = random.choice([
            "Manager",
            "Analyst",
            "Engineer"
        ])


    elif scenario == "CPII-08":

        r["feedback"] = (
            f"Customer {fake.name()} "
            f"contacted support."
        )


    # --------------------------------------------------------
    # FREE TEXT
    # --------------------------------------------------------

    elif scenario == "FREE_TEXT":

        r["feedback"] = (
            f"Please contact {fake.name()} at "
            f"{fake.email()} or "
            f"{fake.numerify('07#########')}."
        )


    # --------------------------------------------------------
    # MULTIPLE DIRECT
    # --------------------------------------------------------

    elif scenario == "MULTIPLE_DIRECT":

        r["customer_name"] = fake.name()

        r["email"] = fake.email()

        r["phone"] = fake.numerify(
            "07#########"
        )

        r["address"] = fake.address().replace(
            "\n",
            ", "
        )


    # --------------------------------------------------------
    # MULTIPLE SENSITIVE
    # --------------------------------------------------------

    elif scenario == "MULTIPLE_SENSITIVE":

        r["medical_condition"] = random.choice([
            "Diabetes",
            "Asthma"
        ])

        r["ethnicity"] = random.choice([
            "Asian",
            "Black",
            "White"
        ])

        r["religion"] = random.choice([
            "Christian",
            "Hindu",
            "Muslim"
        ])


    # --------------------------------------------------------
    # MIXED
    # --------------------------------------------------------

    elif scenario == "MIXED":

        r["customer_name"] = fake.name()

        r["email"] = fake.email()

        r["medical_condition"] = random.choice([
            "Diabetes",
            "Asthma"
        ])

        r["feedback"] = (
            f"Customer {fake.name()} "
            f"contacted support."
        )


    return r


# ============================================================
# RANDOM DISTRIBUTION
# ============================================================

# PASS is randomly selected between 1% and 15%.
pass_count = random.randint(1, 15)


# Choose BLOCK percentage randomly.
#
# We keep it reasonably varied instead of allowing extreme
# datasets where almost everything is BLOCK.
block_count = random.randint(20, 55)


# Remaining records become FLAG.
flag_count = N - pass_count - block_count


# Safety check.
if flag_count < 1:

    block_count = N - pass_count - 1

    flag_count = 1


# ============================================================
# BUILD SCENARIO LIST
# ============================================================

selected = []


# ------------------------------------------------------------
# PASS RECORDS
# ------------------------------------------------------------

selected.extend(
    ["PASS"] * pass_count
)


# ------------------------------------------------------------
# BLOCK RECORDS
# ------------------------------------------------------------

# Guarantee that every BLOCK scenario appears at least once
# whenever there are enough BLOCK records.

block_required = BLOCK_SCENARIOS.copy()

random.shuffle(block_required)

if block_count >= len(block_required):

    selected.extend(block_required)

    remaining_blocks = (
        block_count - len(block_required)
    )

    selected.extend(
        random.choices(
            BLOCK_SCENARIOS,
            k=remaining_blocks
        )
    )

else:

    selected.extend(
        random.sample(
            BLOCK_SCENARIOS,
            block_count
        )
    )


# ------------------------------------------------------------
# FLAG RECORDS
# ------------------------------------------------------------

# Guarantee every FLAG scenario when possible.

flag_required = FLAG_SCENARIOS.copy()

random.shuffle(flag_required)

if flag_count >= len(flag_required):

    selected.extend(flag_required)

    remaining_flags = (
        flag_count - len(flag_required)
    )

    selected.extend(
        random.choices(
            FLAG_SCENARIOS,
            k=remaining_flags
        )
    )

else:

    selected.extend(
        random.sample(
            FLAG_SCENARIOS,
            flag_count
        )
    )


# ============================================================
# FINAL SHUFFLE
# ============================================================

random.shuffle(selected)


# ============================================================
# GENERATE DATASET
# ============================================================

data = [
    generate_record(i, scenario)
    for i, scenario in enumerate(selected, 1)
]


df = pd.DataFrame(
    data,
    columns=COLUMNS
)


# ============================================================
# WRITE EXCEL
# ============================================================

df.to_excel(
    OUTPUT_FILE,
    index=False,
    sheet_name="Synthetic Data"
)


# ============================================================
# FORMAT EXCEL
# ============================================================

wb = load_workbook(OUTPUT_FILE)

ws = wb["Synthetic Data"]

ws.freeze_panes = "A2"

ws.auto_filter.ref = ws.dimensions


# Header formatting

for cell in ws[1]:

    cell.font = Font(bold=True)

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# Column widths

widths = [
    12,
    22,
    30,
    18,
    35,
    15,
    12,
    20,
    18,
    22,
    18,
    22,
    16,
    16,
    20,
    16,
    16,
    18,
    22,
    18,
    45
]


for i, width in enumerate(widths, 1):

    ws.column_dimensions[
        chr(64 + i)
    ].width = width


# Cell formatting

for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


ws.row_dimensions[1].height = 30


wb.save(OUTPUT_FILE)


# ============================================================
# OUTPUT SUMMARY
# ============================================================

print("=" * 55)
print("SYNTHETIC DATA GENERATED")
print("=" * 55)

print(f"\nGenerated {N} records")
print(f"Output: {OUTPUT_FILE}")

print("\nTarget distribution:")
print(f"PASS  : {pass_count} records ({pass_count}%)")
print(f"FLAG  : {flag_count} records ({flag_count}%)")
print(f"BLOCK : {block_count} records ({block_count}%)")

print("\nScenario coverage:")
print(
    pd.Series(selected)
    .value_counts()
    .sort_index()
)

print("\nDone.")